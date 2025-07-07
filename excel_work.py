# import pandas as pd
# import datetime as dt
# from glob import glob
# import logging
# import os

# logging.basicConfig(
#     level=logging.INFO,
#     format='%(asctime)s - %(levelname)s - %(message)s',
#     handlers=[
#         logging.FileHandler("логи_выполнения.log", encoding='utf-8', mode='w'),
#         logging.StreamHandler()
#     ]
# )


# class File:
#     def __init__(self, dataframe, filename):
#         self.df = dataframe
#         self.filename = filename


# TIME_TODAY = dt.date.today()
# WORK_DIR = 'files'
# RESULT_DIR = f'files_result_{TIME_TODAY.strftime('%m.%d.%Y')}'

# FILENAMES = [
#     'Ежедневный финансовый отчет',
#     'Ежедневный_отчет_Склад_*',
#     'Поставщики комплектующих'
# ]


# def load_daily_files():
#     daily_finance_path = glob(f'{WORK_DIR}/Ежедневный финансовый отчет.xlsx')
#     daily_finance_production_path = glob(f'{WORK_DIR}/Ежедневный_отчет_Склад_*')
    
    
    
#     if not daily_finance_path:
#         logging.error('Не найден файл: Ежедневный финансовый отчет.xlsx')
#     if not daily_finance_production_path:
#         logging.error('Не найден файл: Ежедневный_отчет_Склад_...')
        

#     daily_finance = File(
#         pd.read_excel(daily_finance_path[0], sheet_name=None),
#         os.path.basename(daily_finance_path[0])
#     )
#     daily_finance_production = File(
#         pd.read_excel(daily_finance_production_path[0], sheet_name=None),
#         os.path.basename(daily_finance_production_path[0])
#     )
    
    
#     return daily_finance, daily_finance_production


# def result_add(*data_files):
#     try:
#         os.mkdir(RESULT_DIR)
#     except FileExistsError:
#         logging.info(f'Папка {RESULT_DIR} уже создана. Повторное создание не нужно')
        
#     for file in data_files:
#         with pd.ExcelWriter(f'{RESULT_DIR}/{file.filename}') as writer:
#             for sheet_name, sheet_df in file.df.items():
#                 sheet_df.to_excel(writer, sheet_name=sheet_name, index=False)
        
        

# def main():
#     daily_finance, daily_finance_production = load_daily_files()
#     print(type(daily_finance), daily_finance_production)
#     print(dt.date.today())
#     print(dt.datetime.now())
#     result_add(daily_finance, daily_finance_production)
    



# if __name__ == '__main__':
#     main()


from openpyxl import load_workbook, Workbook
from openpyxl.cell.cell import MergedCell
from copy import copy

src_wb = load_workbook("files//Ежедневный финансовый отчет.xlsx")

dest_wb = Workbook()
# Удаляем стандартный лист, т.к. создадим свои
default_sheet = dest_wb.active
dest_wb.remove(default_sheet)

for sheet_name in src_wb.sheetnames:
    src_ws = src_wb[sheet_name]
    dest_ws = dest_wb.create_sheet(title=sheet_name)

    # Копируем значения и стили
    for row in src_ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue  # пропускаем ячейки-части объединения
                
            new_cell = dest_ws.cell(row=cell.row, column=cell.col_idx, value=cell.value)

            if cell.has_style:
                new_cell.font = copy(cell.font)
                new_cell.border = copy(cell.border)
                new_cell.fill = copy(cell.fill)
                new_cell.number_format = copy(cell.number_format)
                new_cell.protection = copy(cell.protection)
                new_cell.alignment = copy(cell.alignment)

    # Копируем ширину столбцов
    for col in src_ws.column_dimensions:
        dest_ws.column_dimensions[col].width = src_ws.column_dimensions[col].width

    # Копируем объединения ячеек
    for merged_range in src_ws.merged_cells.ranges:
        dest_ws.merge_cells(str(merged_range))

dest_wb.save("копия.xlsx")


